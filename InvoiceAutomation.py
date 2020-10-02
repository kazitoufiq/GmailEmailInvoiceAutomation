
from docx import Document
from docx.shared import Inches
from datetime import datetime
from datetime import timedelta
document = Document('Invoice.docx')

InputRentalStart = '2020-10-04'



tables = document.tables



RentalStartYear =datetime.strptime(InputRentalStart, '%Y-%m-%d').strftime("%Y")
RentalStartMonth =datetime.strptime(InputRentalStart, '%Y-%m-%d').strftime("%m")
RentalStartDay =datetime.strptime(InputRentalStart, '%Y-%m-%d').strftime("%d")

tables[0].cell(0,0).text = 'Invoice No_ '+ RentalStartYear + RentalStartMonth + RentalStartDay

NewInvoiceNo = tables[0].cell(0,0).text


# In[55]:


tables[0].cell(0,1).text = "Issued: " +  datetime.today().strftime('%a, %d-%b-%Y')

tables[0].cell(0,1).text


# In[56]:


DueBy ="Due by: " + (datetime.today() + timedelta(days=2)).strftime('%a, %d-%b-%Y')
    
tables[0].cell(0,2).text = DueBy 

DueBy


# In[57]:


tables[1].cell(1,0).text


# In[58]:


tables[2].cell(0,1).text


# In[59]:


tables[3].cell(0,2).text


# In[60]:


tables[3].cell(0,3).text
tables[3].cell(0,4).text

tables[3].cell(0,4).text 


# In[61]:




StartDate=datetime.strptime(InputRentalStart, '%Y-%m-%d').strftime('%A, %d-%b-%Y')


# In[62]:


EndDate=(datetime.strptime(InputRentalStart, '%Y-%m-%d') + 
         timedelta(days=6)).strftime('%A, %d-%b-%Y')


# In[63]:


StartDate


# In[64]:


EndDate


# In[65]:


RentalPeriod = 'Van (Reg: 1OC9KC) Rental Period: '+ StartDate + ' to ' + EndDate +' (7 Days)'


# In[66]:


RentalPeriod


# In[67]:


tables[2].cell(0,1).text = RentalPeriod


# In[68]:




output_new_invoice = NewInvoiceNo +'.docx'

document.save(output_new_invoice)



from docx2pdf import convert

new_invoice_pdf_file = NewInvoiceNo+'.pdf'

convert(output_new_invoice,  new_invoice_pdf_file)